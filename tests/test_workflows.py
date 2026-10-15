import base64
from io import BytesIO
import json
import pytest

import storage as s
from importers import extract

@pytest.fixture
def fresh_db(tmp_path,monkeypatch):
    monkeypatch.setenv('FRIDAY_NIGHT_IQ_DB',str(tmp_path/'test.sqlite'))
    s.init_db()
    return s.rows('SELECT * FROM teams')[0]['id']


def test_demo_team_starts_with_usable_play_cards(fresh_db):
    assert len(s.rows('SELECT * FROM plays WHERE team_id=?',(fresh_db,)))==3
    assert len(s.rows('SELECT * FROM players WHERE team_id=?',(fresh_db,)))==4


def test_position_filter_and_study_progress(fresh_db):
    team=fresh_db
    qb=s.one("SELECT id FROM players WHERE team_id=? AND position='QB'",(team,))['id']
    ol=s.one("SELECT id FROM players WHERE team_id=? AND position='OL'",(team,))['id']
    all_plays=s.rows('SELECT * FROM plays WHERE team_id=?',(team,))
    install=s.create_install(team,'Base','Test Opponent','Week 1','2027-09-01','', [p['id'] for p in all_plays])
    qb_visible=s.player_install_plays(team,qb,install)
    ol_visible=s.player_install_plays(team,ol,install)
    assert len(qb_visible)==2
    assert len(ol_visible)==1
    target=next(p['id'] for p in qb_visible if p['title']=='Mesh Concept')
    s.set_progress(team,qb,install,target,'Ready')
    status={p['id']:p['study_status'] for p in s.player_install_plays(team,qb,install)}
    assert status[target]=='Ready'
    with pytest.raises(ValueError):
        s.set_progress(team,ol,install,target,'Ready')


def test_team_isolation(fresh_db):
    second=s.create_team('Rival','2027')
    team=fresh_db
    foreign_play=s.rows('SELECT id FROM plays WHERE team_id=?',(team,))[0]['id']
    pid=s.write("INSERT INTO players(team_id,name,position) VALUES (?,?,?)",(second,'Second Team Player','QB'))
    install=s.create_install(second,'First','Other','1','2027-08-01','',[foreign_play])
    assert s.rows('SELECT * FROM install_plays WHERE install_id=?',(install,))==[]
    assert s.player_install_plays(second,pid,install)==[]
    with pytest.raises(ValueError):
        s.set_progress(second,pid,install,foreign_play,'Ready')


def test_export_includes_complete_data(fresh_db):
    doc=s.write("INSERT INTO uploads(team_id,filename,filetype,filebytes,segments_json,warnings_json) VALUES(?,?,?,?,?,?)",
                (fresh_db,'clip.png','png',b'\x89PNGsample','[]','[]'))
    data=json.loads(s.export_team(fresh_db))
    assert data['format']=='friday-night-iq-backup-v1'
    assert data['uploads'][0]['id']==doc
    assert base64.b64decode(data['uploads'][0]['filebytes'])==b'\x89PNGsample'


def test_text_and_csv_extraction():
    segments,warnings=extract('offense.txt',b'Outside zone\nBlock right\n---\nCounter GT\nPull left')
    assert len(segments)==2
    assert segments[1]['heading']=='Counter GT'
    assert not warnings
    rows,_=extract('plays.csv',b'name,assignment\nMesh,Run shallow cross\nZone,Block play side\n')
    assert len(rows)==2


def test_image_keeps_attachment_without_fabricated_text():
    parts,warnings=extract('play.png',b'fake-image-for-upload-test')
    assert parts==[]
    assert 'Image stored' in warnings[0]
    with pytest.raises(ValueError): extract('virus.exe',b'x')
    with pytest.raises(ValueError): extract('blank.pdf',b'')


def test_docx_extraction():
    from docx import Document
    doc=Document();doc.add_heading('Inside Zone',level=1);doc.add_paragraph('Covered step, vertical displacement')
    doc.add_heading('Power',level=1);doc.add_paragraph('Pull the backside guard')
    b=BytesIO();doc.save(b)
    parts,_=extract('playbook.docx',b.getvalue())
    assert len(parts)==2
    assert parts[1]['heading']=='Power'


def test_pptx_extraction():
    from pptx import Presentation
    p=Presentation(); slide=p.slides.add_slide(p.slide_layouts[1]);slide.shapes.title.text='Trips Mesh'
    b=BytesIO();p.save(b)
    parts,_=extract('playbook.pptx',b.getvalue())
    assert len(parts)==1
    assert parts[0]['heading']=='Trips Mesh'

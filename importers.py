"""Extract *candidate* play sections from uploads; never claim assignments were inferred.

Diagrams embedded as images stay attached in the original file. Scanned PDFs
are retained without OCR, and will need manual play entry or future OCR.
"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path
import csv
import re

SUPPORTED={"pdf", "docx", "pptx", "txt", "md", "csv", "png", "jpg", "jpeg"}
MAX_UPLOAD_MB=15


def clean(s: str) -> str:
    return re.sub(r"[ \t]+", " ", str(s or "")).strip()


def title_from(text: str, fallback: str) -> str:
    for line in text.splitlines():
        line=clean(line).strip("•–-: ")
        if 3 <= len(line) <= 70 and len(line.split()) <= 12:
            return line
    return fallback


def result_segment(text:str, fallback:str, ordinal:int):
    text=clean(text).strip()
    return {"heading":title_from(text,fallback),"text":text[:20000],"source":fallback,"ordinal":ordinal}


def extract(filename:str, content:bytes):
    """Return (segments, warnings); segments are coach-review drafts only."""
    suffix=Path(filename).suffix.lower().lstrip(".")
    if suffix not in SUPPORTED: raise ValueError("Unsupported extension")
    if not content: raise ValueError("Empty file")
    if len(content) > MAX_UPLOAD_MB * 1024 * 1024: raise ValueError(f"File must be under {MAX_UPLOAD_MB} MB")
    warnings=[]
    parts=[]
    try:
        if suffix=="pdf":
            from pypdf import PdfReader
            pdf=PdfReader(BytesIO(content))
            if pdf.is_encrypted:
                raise ValueError("Encrypted PDFs are not supported")
            for i,page in enumerate(pdf.pages,1):
                txt=page.extract_text() or ""
                if txt.strip(): parts.append(result_segment(txt,f"Page {i}",i))
            if len(parts)<len(pdf.pages):
                warnings.append("Some PDF pages contain no extractable text. Diagram-only/scanned pages are preserved in the uploaded file, but aren't auto-converted into plays.")
        elif suffix=="pptx":
            from pptx import Presentation
            pres=Presentation(BytesIO(content))
            for i,slide in enumerate(pres.slides,1):
                txt="\n".join(shape.text for shape in slide.shapes if shape.has_text_frame and shape.text.strip())
                if txt.strip(): parts.append(result_segment(txt,f"Slide {i}",i))
            if len(parts)<len(pres.slides):
                warnings.append("Slides containing only diagrams/images won't have extractable text; original slides are preserved.")
        elif suffix=="docx":
            from docx import Document
            doc=Document(BytesIO(content))
            current=[]
            name="Section 1"
            section=1
            for p in doc.paragraphs:
                if not p.text.strip(): continue
                if p.style and p.style.name and p.style.name.lower().startswith("heading") and current:
                    parts.append(result_segment("\n".join(current),name,section))
                    section+=1
                    current=[]
                    name=p.text.strip()
                current.append(p.text)
            # Include table text rather than silently dropping all tables.
            for table in doc.tables:
                for row in table.rows:
                    current.append(" | ".join(cell.text for cell in row.cells))
            if current: parts.append(result_segment("\n".join(current),name,section))
        elif suffix in {"txt","md","csv"}:
            txt=content.decode("utf-8-sig",errors="replace")
            if suffix=="csv":
                reader=csv.DictReader(BytesIO(content).read().decode("utf-8-sig",errors="replace").splitlines())
                for i,line in enumerate(reader,1):
                    desc="\n".join(f"{k}: {v}" for k,v in line.items() if k and v)
                    if desc.strip(): parts.append(result_segment(desc,f"Row {i}",i))
            else:
                # Delimiters explicitly supplied by the coach (---) are reliable.
                chunks=[p.strip() for p in re.split(r"(?m)^\s*---+\s*$",txt) if p.strip()]
                for i,c in enumerate(chunks,1): parts.append(result_segment(c,f"Section {i}",i))
        else:
            warnings.append("Image stored as an attachment. Text recognition and diagram-to-play conversion aren't available in this prototype; create the play details manually.")
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"Could not read the {suffix.upper()} file. Check that it is valid and not password-protected.") from exc
    if not parts:
        warnings.append("No readable text sections found. You can still make a manual play and associate it with the uploaded source.")
    if len(parts)>100:
        parts=parts[:100]
        warnings.append("Showing the first 100 sections to keep the demo responsive.")
    return parts,warnings
